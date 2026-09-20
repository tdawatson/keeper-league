#!/usr/bin/env python3
"""Import third-party projections and score them with THIS league's rules.

Currently supports the Apples & Ginos public workbook (skaters). Their sheet is
published openly at applesandginos.com/skater-projections/ as a Google Sheet;
fetch it with --refresh, or work from the saved copy in data/.

Writes data/proj-external.csv: one row per matched player, per source.
"""
import argparse
import csv
import json
import pathlib
import re
import unicodedata
import urllib.request

AG_SHEET = "17XHuu3nWDYMGGz-xUtEVt8QJTjTEoQEp4-qZJnnVJq0"
AG_LOCAL = "data/applesandginos-2627.xlsx"
AG_TABS = {"A&G Avg": "A&G Avg Projections",
           "A&G Nate": "Nates Projections",
           "A&G Blake": "Blakes Projections"}

# A&G column layout (0-based), read off the header row
COLS = {"name": 0, "team": 1, "pos": 2, "GP": 4, "G": 5, "A": 6,
        "SOG": 9, "HIT": 10, "BLK": 11, "PIM": 12}


def norm(n):
    """Normalise a player name for cross-source matching."""
    if not n:
        return ""
    n = unicodedata.normalize("NFKD", str(n))
    n = "".join(c for c in n if not unicodedata.combining(c))
    n = n.lower().replace("&", "and")
    n = re.sub(r"\b(jr|sr|ii|iii|iv)\b", "", n)
    n = re.sub(r"[^a-z ]", "", n)
    return re.sub(r"\s+", " ", n).strip()


def scoring_fn(info):
    """Build fpts(stats, is_d) from the league's own scoring categories."""
    cats = info["scoringSystem"]["scoringCategories"]["SKATING"]

    def rate(cat, is_d):
        v = cats.get(cat)
        if not v:
            return 0.0
        s = v.get("D" if is_d else "Default", v.get("Default", "points0"))
        return float(s.replace("points", "")) if s.startswith("points") else 0.0

    def fpts(s, is_d):
        return round(sum(s.get(c, 0.0) * rate(c, is_d)
                         for c in ("G", "A", "PIM", "Blk", "Hit", "SOG")), 2)

    # our stat keys -> league category codes
    return lambda s, is_d: fpts({"G": s["G"], "A": s["A"], "PIM": s["PIM"],
                                 "Blk": s["BLK"], "Hit": s["HIT"], "SOG": s["SOG"]}, is_d)


def read_ag(path, tab):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)[tab]
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows)
                 if r and str(r[0]).strip() == "Name") + 1
    out = []
    for r in rows[start:]:
        if not r or not r[COLS["name"]]:
            continue
        rec = {"name": str(r[COLS["name"]]).strip(),
               "team": r[COLS["team"]], "pos": r[COLS["pos"]]}
        ok = True
        for k in ("GP", "G", "A", "SOG", "HIT", "BLK", "PIM"):
            v = r[COLS[k]]
            if not isinstance(v, (int, float)):
                ok = False
                break
            rec[k] = float(v)
        if ok:
            out.append(rec)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", action="store_true",
                    help="re-download the A&G sheet before importing")
    ap.add_argument("--out", default="data/proj-external.csv")
    args = ap.parse_args()

    if args.refresh:
        url = f"https://docs.google.com/spreadsheets/d/{AG_SHEET}/export?format=xlsx"
        req = urllib.request.Request(url, headers={"User-Agent": "keeper-league/0.1"})
        with urllib.request.urlopen(req, timeout=60) as r:
            pathlib.Path(AG_LOCAL).write_bytes(r.read())
        print(f"refreshed {AG_LOCAL}")

    d = pathlib.Path("data")
    info = json.loads((d / "getLeagueInfo.json").read_text())
    pinfo = info["playerInfo"]
    fx = {r["ID"].strip("*"): r
          for r in csv.DictReader((d / "fantrax-actual-2526.csv").open())}
    fpts = scoring_fn(info)

    # Match on name + NHL team + F/D/G. All three are needed: Vancouver
    # carries two Elias Petterssons, one forward and one defenceman, so name
    # and team alone silently picks the wrong one and scores a forward's
    # stat line with defenceman multipliers.
    def bucket(pos):
        p = (pos or "").strip().upper()
        return "G" if p.startswith("G") else ("D" if p.startswith("D") else "F")

    exact, by_team, by_name, loose = {}, {}, {}, {}
    for pid, r in fx.items():
        k = norm(r["Player"])
        t = (r["Team"] or "").strip().upper()
        b = bucket(pinfo.get(pid, {}).get("eligiblePos") or r["Position"])
        exact.setdefault((k, t, b), []).append(pid)
        by_team.setdefault((k, t), []).append(pid)
        by_name.setdefault(k, []).append(pid)
        parts = k.split()
        if len(parts) >= 2:
            loose.setdefault((parts[0][:1], parts[-1], t, b), []).append(pid)

    ambiguous = []

    def lookup(name, team, pos):
        k, t, b = norm(name), (team or "").strip().upper(), bucket(pos)
        for table, key in ((exact, (k, t, b)), (by_team, (k, t)),
                           (by_name, k), (loose, (k.split()[0][:1], k.split()[-1], t, b))):
            cand = table.get(key, [])
            if len(cand) == 1:
                return cand[0]
            if len(cand) > 1 and table is by_name:
                ambiguous.append(name)
                return None
        return None

    rows, report = [], []
    for label, tab in AG_TABS.items():
        recs = read_ag(AG_LOCAL, tab)
        hit = miss = 0
        unmatched = []
        for r in recs:
            pid = lookup(r["name"], r["team"], r["pos"])
            if not pid:
                miss += 1
                unmatched.append(r["name"])
                continue
            hit += 1
            elig = pinfo.get(pid, {}).get("eligiblePos", "")
            is_d = elig == "D"
            rows.append({"source": label, "fantraxId": pid,
                         "player": fx[pid]["Player"], "pos": elig or r["pos"],
                         "GP": r["GP"], "G": r["G"], "A": r["A"], "SOG": r["SOG"],
                         "HIT": r["HIT"], "BLK": r["BLK"], "PIM": r["PIM"],
                         "FPts": fpts(r, is_d)})
        report.append((label, len(recs), hit, miss, unmatched[:8]))

    with open(args.out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)

    print(f"wrote {args.out}  ({len(rows)} rows)\n")
    print(f"{'source':<12}{'rows':>7}{'matched':>9}{'missed':>8}   sample unmatched")
    for label, n, hit, miss, un in report:
        print(f"{label:<12}{n:>7}{hit:>9}{miss:>8}   {', '.join(un[:4])}")
    if ambiguous:
        print(f"\nnames too ambiguous to match safely: {len(set(ambiguous))} "
              f"({', '.join(sorted(set(ambiguous))[:4])})")


if __name__ == "__main__":
    main()
